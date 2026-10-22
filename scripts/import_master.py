from __future__ import annotations
import argparse, json, re
from pathlib import Path
from openpyxl import load_workbook


def clean(v):
    if v is None: return None
    if isinstance(v, str):
        v=v.strip()
        return v or None
    return str(v).strip()


def sql_text(v):
    if v is None: return 'null'
    return "'" + str(v).replace("'", "''") + "'"


def page(v):
    if v in (None,''): return 'null'
    return str(int(v))


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('xlsx', type=Path)
    ap.add_argument('--out', type=Path, required=True)
    args=ap.parse_args()
    wb=load_workbook(args.xlsx, data_only=True, read_only=True)
    required={'Knoten','Verknuepfungen','Abgleich_OZ','Hinweise'}
    if set(wb.sheetnames) != required:
        raise SystemExit(f'Unexpected worksheets: {wb.sheetnames}')

    def records(name):
        ws=wb[name]; rows=ws.iter_rows(values_only=True); h=[clean(x) for x in next(rows)]
        return [dict(zip(h,r)) for r in rows if any(x is not None for x in r)]

    nodes=records('Knoten'); edges=records('Verknuepfungen'); audit=records('Abgleich_OZ')
    ids=[clean(x['id']) for x in nodes]
    if len(ids)!=len(set(ids)): raise SystemExit('Duplicate node IDs')
    if any(clean(x['verbindungstyp'])!='Hierarchie' for x in edges): raise SystemExit('Unexpected non-hierarchy relation in master')
    idset=set(ids)
    broken=[e for e in edges if clean(e['quelle_id']) not in idset or clean(e['ziel_id']) not in idset]
    if broken: raise SystemExit(f'Broken edge references: {broken[:3]}')
    key_nodes={(clean(n['dezernat']),clean(n['oz'])) for n in nodes if clean(n['typ']) in ('Themenbereich','Aufgabe')}
    key_audit={(clean(a['dezernat']),clean(a['oz'])) for a in audit}
    if key_nodes != key_audit: raise SystemExit('Abgleich_OZ differs from Knoten OZ set')

    parents={clean(e['ziel_id']):clean(e['quelle_id']) for e in edges}
    if len(parents)!=len(edges): raise SystemExit('A formal child has multiple hierarchy parents')

    dez={clean(n['id']):n for n in nodes if clean(n['typ'])=='Dezernat'}
    sg={clean(n['id']):n for n in nodes if clean(n['typ'])=='Sachgebiet'}
    topics={clean(n['id']):n for n in nodes if clean(n['typ'])=='Themenbereich'}
    tasks={clean(n['id']):n for n in nodes if clean(n['typ'])=='Aufgabe'}

    def find_sg(node_id):
        cur=node_id; seen=set()
        while cur in parents:
            cur=parents[cur]
            if cur in seen: raise SystemExit('Hierarchy cycle')
            seen.add(cur)
            if cur in sg: return cur
        raise SystemExit(f'No subject area ancestor for {node_id}')

    summary={'nodes':len(nodes),'departments':len(dez),'subject_areas':len(sg),'topics':len(topics),'tasks':len(tasks),'hierarchy_edges':len(edges),'oz_audit':len(audit)}
    lines=['begin;','',"insert into public.datasets (id,name,effective_date,source_file,source_reference,active) values ('00000000-0000-0000-0000-000000000001','GVP 2026-09-01','2026-09-01','GVP_Masterdaten_final.xlsx','Nur-GB_3_2026-09-01_GVPL_ZGB.pdf',true);",'']
    DS="'00000000-0000-0000-0000-000000000001'"

    for group, typ in ((dez,'dezernat'),(sg,'sachgebiet')):
        for xid,n in group.items():
            name=clean(n['originalbezeichnung']) or clean(n['anzeigebezeichnung'])
            lines.append(f"insert into public.organization_units (dataset_id,external_id,unit_type,name,short_name,internal) values ({DS},{sql_text(xid)},'{typ}',{sql_text(name)},{sql_text(clean(n['anzeigebezeichnung']))},true);")
    for ext,name in [('EXT-GB1','anderer Geschäftsbereich 1'),('EXT-GB2','anderer Geschäftsbereich 2'),('EXT-GB4','anderer Geschäftsbereich 4')]:
        lines.append(f"insert into public.organization_units (dataset_id,external_id,unit_type,name,short_name,internal) values (null,{sql_text(ext)},'geschaeftsbereich',{sql_text(name)},{sql_text(name)},false);")
    lines.append('')

    for xid,n in topics.items():
        lines.append("insert into public.topics (dataset_id,external_id,department_no,oz,original_name,display_name,source,pointer_raw,pdf_page) values " +
                     f"({DS},{sql_text(xid)},{sql_text(clean(n['dezernat']))},{sql_text(clean(n['oz']))},{sql_text(clean(n['originalbezeichnung']))},{sql_text(clean(n['anzeigebezeichnung']))},{sql_text(clean(n['quelle']) or 'GVP')},{sql_text(clean(n['weiserzeichen']))},{page(n['pdf_seite'])});")
    lines.append('')

    for xid,n in tasks.items():
        sgid=find_sg(xid); dept='D-'+clean(n['dezernat']); p=parents.get(xid); topic=p if p in topics else None
        lines.append("insert into public.tasks (dataset_id,external_id,source,department_id,subject_area_id,topic_id,oz,original_name,display_name,pointer_raw,pdf_page) select " +
                     f"{DS},{sql_text(xid)},'GVP',d.id,s.id,t.id,{sql_text(clean(n['oz']))},{sql_text(clean(n['originalbezeichnung']))},{sql_text(clean(n['anzeigebezeichnung']))},{sql_text(clean(n['weiserzeichen']))},{page(n['pdf_seite'])} " +
                     f"from public.organization_units d join public.organization_units s on s.external_id={sql_text(sgid)} left join public.topics t on t.external_id={sql_text(topic)} and t.dataset_id={DS} where d.external_id={sql_text(dept)} and d.dataset_id={DS};")
    lines.append('')

    kinds={**{x:'organization_unit' for x in dez|sg}, **{x:'topic' for x in topics}, **{x:'task' for x in tasks}}
    for e in edges:
        a,b=clean(e['quelle_id']),clean(e['ziel_id'])
        lines.append("insert into public.hierarchy_edges (dataset_id,source_external_id,target_external_id,source_kind,target_kind,provenance,note) values " +
                     f"({DS},{sql_text(a)},{sql_text(b)},{sql_text(kinds[a])},{sql_text(kinds[b])},{sql_text(clean(e['herkunft']) or 'GVP')},{sql_text(clean(e['hinweis']))});")
    lines += ['', 'commit;']
    args.out.write_text('\n'.join(lines)+'\n', encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False))

if __name__=='__main__': main()
